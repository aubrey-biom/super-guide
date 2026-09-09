"""Guards on the B&M master-forecast parser, and the combine's own self-check.

Every fixture is built in-process with openpyxl, so there is no committed slice of the
sheet to go stale (the retired owner parser's fixtures did, and were deleted with it).
Each test pins ONE guard, and the guard is the reason the test exists: a sheet that
arrives by hand is the input most likely to be filtered, sorted or re-headed between
runs, which is the 2026-07-29 P0 shape.
"""

from __future__ import annotations

import calendar
from datetime import date

import pytest
from openpyxl import Workbook

from pipelines.target_shipment_forecast.inputs.bm_master_forecast import (
    BLOCK_FLOOR,
    METRICS,
    BmParseError,
    parse_target_schedule,
)
from pipelines.target_shipment_forecast.model.bm_combine import demo as combine_demo

MONTHS = [date(2026, m, 1) for m in range(1, 13)] + [date(2027, m, 1) for m in range(1, 13)]


def _sheet(tmp_path, *, blocks=BLOCK_FLOOR, header=("Metric", "SKU / Description", "Unique Key"),
           days_label="Days in Month", days=None, mutate=None):
    """A minimal, structurally valid `Target Schedule`, then whatever `mutate` does to it."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Target Schedule"
    ws.append([])
    ws.append([])
    ws.append(list(header) + [f"{m:%b-%Y}" for m in MONTHS])
    ws.append([days_label, None, None] + (days or [calendar.monthrange(m.year, m.month)[1] for m in MONTHS]))
    for i in range(blocks):
        sku = f"P-SKU-{i:03d}"
        # Vary per block: identical series across blocks is itself a finding the parser
        # reports (BM_IDENTICAL_SERIES), so a valid fixture must not be 30 copies.
        stores = [1500.0 + i] * len(MONTHS)
        upspw = [1.0 + i / 100.0] * len(MONTHS)
        vel = [s * u * calendar.monthrange(m.year, m.month)[1] / 7.0
               for s, u, m in zip(stores, upspw, MONTHS)]
        load = [0.0] * len(MONTHS)
        rows = {
            "Stores": stores, "UPSPW": upspw, "Velocity": vel, "Load_Orders": load,
            "Quote": [3.0] * len(MONTHS),
            "Total_Demand": [v + l for v, l in zip(vel, load)],
            "Revenue": [(v + l) * 3.0 for v, l in zip(vel, load)],
        }
        for k, metric in enumerate(METRICS):
            b = sku if metric == "Stores" else (f"desc {i}" if metric == "UPSPW" else None)
            c = f"TARGET_{sku}" if metric == "Stores" else (sku if metric == "UPSPW" else None)
            ws.append([metric, b, c] + rows[metric])
        ws.append([])
    if mutate is not None:
        mutate(ws)
    p = tmp_path / "bm.xlsx"
    wb.save(p)
    return p


def _parse(path):
    """Parse with resolution switched off -- these tests are about structure, not mapping."""
    class _NoItems:
        def tcins_for(self, *_a, **_k):
            return []

    class _NoAliases:
        def resolve(self, raw, *, kind="auto"):
            class R:
                canonical_sku = str(raw)
            return R()

    return parse_target_schedule(path, item_master=_NoItems(), aliases=_NoAliases())


def test_a_valid_sheet_parses_and_carries_every_metric(tmp_path):
    bm = _parse(_sheet(tmp_path))
    assert len(bm.months) == len(MONTHS)
    assert bm.frame.shape[0] == BLOCK_FLOOR * len(MONTHS)
    for metric in METRICS:
        assert f"bm_{metric.lower()}" in bm.frame.columns
    # snapshot_date is the file's mtime, never today()
    assert bm.snapshot_date <= date.today()
    assert not bm.warnings


def test_a_renamed_header_aborts_rather_than_parsing_by_position(tmp_path):
    p = _sheet(tmp_path, header=("Measure", "SKU / Description", "Unique Key"))
    with pytest.raises(BmParseError, match="no header row"):
        _parse(p)


def test_a_truncated_sheet_aborts_on_the_block_floor(tmp_path):
    p = _sheet(tmp_path, blocks=BLOCK_FLOOR - 1)
    with pytest.raises(BmParseError, match="below floor"):
        _parse(p)


def test_text_in_a_numeric_cell_aborts_and_names_the_cell(tmp_path):
    def mutate(ws):
        ws.cell(row=5, column=4).value = "TBD"  # first block's Stores, Jan-2026

    with pytest.raises(BmParseError, match=r"holds text 'TBD'"):
        _parse(_sheet(tmp_path, mutate=mutate))


def test_a_blank_numeric_cell_is_zero_not_an_error(tmp_path):
    def mutate(ws):
        ws.cell(row=5, column=4).value = None
        ws.cell(row=7, column=4).value = 0  # keep Velocity consistent
        ws.cell(row=10, column=4).value = 0  # Total_Demand
        ws.cell(row=11, column=4).value = 0  # Revenue

    bm = _parse(_sheet(tmp_path, mutate=mutate))
    assert float(bm.frame["bm_stores"].iloc[0]) == 0.0


def test_a_days_in_month_row_that_disagrees_with_the_calendar_aborts(tmp_path):
    bad = [calendar.monthrange(m.year, m.month)[1] for m in MONTHS]
    bad[1] = 31  # February is not 31 days, and it is the Velocity denominator
    with pytest.raises(BmParseError, match="disagrees with the calendar"):
        _parse(_sheet(tmp_path, days=bad))


def test_a_missing_days_in_month_row_aborts(tmp_path):
    with pytest.raises(BmParseError, match="no 'days in month' row"):
        _parse(_sheet(tmp_path, days_label="Calendar Days"))


def test_the_definitional_identity_aborts_but_the_derived_one_only_warns(tmp_path):
    def break_definition(ws):
        ws.cell(row=10, column=4).value = 999999.0  # Total_Demand != Velocity + Load_Orders

    with pytest.raises(BmParseError, match="definitional identity"):
        _parse(_sheet(tmp_path, mutate=break_definition))

    def break_derivation(ws):
        # Velocity hand-overridden; Total_Demand kept consistent with it, so only the
        # derivation Stores x UPSPW x Days/7 fails. This is the live P-60WIP-DSN-ALP
        # Jun-2026 shape and it must not stop the parse.
        ws.cell(row=7, column=4).value = 9999.0
        ws.cell(row=10, column=4).value = 9999.0
        ws.cell(row=11, column=4).value = 9999.0 * 3.0

    bm = _parse(_sheet(tmp_path, mutate=break_derivation))
    assert any(w.startswith("BM_VELOCITY_IDENTITY") for w in bm.warnings), bm.warnings


def test_a_duplicate_unique_key_aborts_with_both_rows(tmp_path):
    def mutate(ws):
        ws.cell(row=13, column=3).value = "TARGET_P-SKU-000"  # second block's key

    with pytest.raises(BmParseError, match="appears twice, at rows"):
        _parse(_sheet(tmp_path, mutate=mutate))


def test_a_placeholder_block_is_flagged_not_dropped(tmp_path):
    def mutate(ws):
        for col in range(4, 4 + len(MONTHS)):
            ws.cell(row=5, column=col).value = 1.0  # first block: one token door

    bm = _parse(_sheet(tmp_path, mutate=mutate))
    assert bm.placeholder_skus == ["P-SKU-000"]
    assert any("BM_PLACEHOLDER_BLOCK" in w for w in bm.warnings)
    # flagged, and still present in the frame
    assert (bm.frame["bm_sku"] == "P-SKU-000").any()


def test_a_missing_metric_row_aborts_and_names_the_block(tmp_path):
    def mutate(ws):
        ws.delete_rows(8)  # first block's Load_Orders row

    with pytest.raises(BmParseError, match="missing metric row"):
        _parse(_sheet(tmp_path, mutate=mutate))


def test_a_missing_tab_aborts(tmp_path):
    wb = Workbook()
    wb.active.title = "Not It"
    p = tmp_path / "x.xlsx"
    wb.save(p)
    with pytest.raises(BmParseError, match="not in"):
        _parse(p)


def test_bm_combine_self_check():
    """The ramp refusals, the fact-stream rule, the grade penalty and the load verdicts."""
    combine_demo()
