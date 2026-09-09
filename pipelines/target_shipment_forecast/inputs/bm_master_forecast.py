"""Brick & Mortar Master Forecast, `Target Schedule` tab -> a tidy TCIN x month frame.

WHAT THIS FILE IS FOR. `dist_velocity` (model/consumption.py) measures Target demand
from BPD and beats every alternative on WAPE, but it cannot see two things, and both
were measured live on 2026-09-09:

  1. FORWARD DISTRIBUTION. `dly_po_plan_tcin.ORDER_D` reaches 2026-11-06 and every
     other BPD feed is historical, so beyond ~2 months `dist_velocity` holds its store
     count flat at the item's own peak. 30 of the 39 SKUs in this sheet have a store
     plan that MOVES after Nov-2026 (chain 24,585 door-slots Sep-26 -> 46,073 Dec-28).
     That plan is Target's/BIOM's stated rollout intent and exists nowhere in BigQuery.
  2. LAUNCH LOADS BEYOND THE PLAN HORIZON. `Load_Orders` is the pipeline-fill row.

WHAT IT IS NOT FOR. This sheet's `Stores` row is NOT a substitute for a measured store
count and this parser does not let it become one (see model/bm_combine.py). Measured
2026-09-09 against the live feeds: `Stores` sits flat at 1,575 for every mature item
while `weekly_inv_tcin_loc.ITEM_LOCATIONS_BASE_COUNT` reads 1,626-1,691, and it is
plainly wrong on five live items -- K-DIS-2BAB-PUR/LGR say 1 store against 1,411/1,370
stocked, P-DIS-LGR/BLK/DGR say 1 store at UPSPW 12.0 (the placeholder shape this parser
flags), and P-20WIP-SAN-STL-TRV says 0 stores / 0 UPSPW while selling 11,459 units in
eight weeks. So the sheet supplies SHAPE, BPD supplies LEVEL.

GUARD DISCIPLINE, following scripts/load_cost_files.py and
pipelines/rdz_inventory/rdz_inventory_ingestion.py:

  - the tab is found by NAME; a missing tab aborts and nothing is returned
  - the header row is found by its header TEXT, never by position, and the month
    columns are parsed from that row's labels
  - `Days in Month` must equal the real calendar length of every month it labels,
    or the whole parse aborts: it is the denominator of the sheet's own identity
  - a row-count FLOOR on SKU blocks (the 2026-07-29 P0 shape: a filtered or sorted
    sheet parsing "successfully" with most rows gone)
  - a block missing any of the seven metric rows aborts, naming the block
  - a non-blank STRING in a numeric cell aborts, naming the block, metric and month.
    Blank is legitimately zero; text is not
  - a duplicate `Unique Key` aborts with both row numbers. Two blocks for one SKU is a
    human disagreement, and picking one or averaging them would hide it -- the rule
    load_cost_files.py applies to conflicting BOM sources
  - `Total_Demand == Velocity + Load_Orders` ABORTS on failure. It is definitional --
    an addition performed inside the sheet -- and it held on 1,872 of 1,872 cells live
  - `Velocity == Stores x UPSPW x Days/7` WARNS, and names the cell. It is a
    DERIVATION the sheet's author may legitimately override by hand, not a definition.
    Live it fails exactly once: P-60WIP-DSN-ALP Jun-2026 carries 15,379.875, which is
    1,575 x 2.205 x 31/7 -- the 31-day divisor of a neighbouring column against June's
    30 days, i.e. a copy-paste in the sheet, +3.3% on that cell
  - `snapshot_date` comes from the file's mtime, never `today()`

SKU -> TCIN resolution reuses what the repo already has, in this order, and never
guesses: `ItemMaster.tcins_for` over `biom_sku`, `rdz_item` and `vendor_style`, with
`data/sku_aliases_target.tsv` applied first. Anything left over stays in `unresolved`
and is reported, never dropped -- the contract the retired owner parser had.
"""

from __future__ import annotations

import calendar
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from pipelines.target_shipment_forecast.inputs.aliases import AliasMap
from pipelines.target_shipment_forecast.inputs.item_master import ItemMaster

SHEET = "Target Schedule"

# The seven metric rows of every SKU block, in the sheet's own order.
METRICS = ("Stores", "UPSPW", "Velocity", "Load_Orders", "Quote", "Total_Demand", "Revenue")

# Header cells that identify the header row, by text.
HEADER_METRIC = "metric"
HEADER_SKU = "sku / description"
HEADER_KEY = "unique key"

DAYS_ROW_LABEL = "days in month"

# Row-count floor on SKU blocks. 39 blocks live on 2026-09-09 (verified by reading the
# file). 30 leaves room for a genuine assortment cut without letting a truncated or
# filtered sheet through.
BLOCK_FLOOR = 30

# Minimum month columns. 48 live (Jan-2025..Dec-2028). 24 is two years -- below that the
# sheet has lost its forward reach, which is the only reason this file is read at all.
MONTH_FLOOR = 24

_MONTH_RE = re.compile(r"^([A-Za-z]{3})-(\d{4})$")


class BmParseError(Exception):
    """The sheet did not verify. Nothing is returned; fix the file."""


@dataclass
class BmForecast:
    """Parsed `Target Schedule`."""

    frame: pd.DataFrame  # bm_sku, unique_key, description, tcin, month_start, <metrics>
    months: list[date]
    snapshot_date: date
    source_file: str
    warnings: list[str] = field(default_factory=list)
    unresolved: list[dict[str, Any]] = field(default_factory=list)
    placeholder_skus: list[str] = field(default_factory=list)
    duplicate_series: list[str] = field(default_factory=list)

    @property
    def tcins(self) -> list[int]:
        """Resolved TCINs, ascending."""
        s = self.frame["tcin"].dropna()
        return sorted({int(t) for t in s})


def _norm(v: object) -> str:
    return " ".join(str("" if v is None else v).strip().lower().split())


def _parse_month(label: object) -> date | None:
    m = _MONTH_RE.match(str("" if label is None else label).strip())
    if not m:
        return None
    try:
        return datetime.strptime(f"{m.group(1)}-{m.group(2)}", "%b-%Y").date()
    except ValueError:
        return None


def _num(v: object, *, where: str) -> float:
    """Blank -> 0.0. A number -> float. Anything else aborts: text in a numeric cell."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return 0.0
    if isinstance(v, bool):
        raise BmParseError(f"ABORT: {where} holds a boolean, not a number")
    if isinstance(v, (int, float)):
        f = float(v)
        if f != f or f in (float("inf"), float("-inf")):
            raise BmParseError(f"ABORT: {where} is not finite ({v!r})")
        return f
    raise BmParseError(f"ABORT: {where} holds text {str(v)[:40]!r}; a numeric cell may be blank, not text")


def _find_header(rows: list[list[Any]]) -> int:
    for i, r in enumerate(rows):
        cells = [_norm(c) for c in r[:3]]
        if cells[:3] == [HEADER_METRIC, HEADER_SKU, HEADER_KEY]:
            return i
    raise BmParseError(
        f"ABORT: no header row in '{SHEET}' whose first three cells are "
        f"{[HEADER_METRIC, HEADER_SKU, HEADER_KEY]!r}. Columns are resolved by name, "
        "never by position, so the parse cannot continue."
    )


def parse_target_schedule(
    path: str | Path,
    *,
    item_master: ItemMaster | None = None,
    aliases: AliasMap | None = None,
    snapshot_date: date | None = None,
) -> BmForecast:
    """Parse and verify the `Target Schedule` tab. Raises `BmParseError` on any
    structural failure; returns a `BmForecast` with warnings otherwise."""
    p = Path(path)
    if not p.exists():
        raise BmParseError(f"ABORT: {p} does not exist")
    snap = snapshot_date or datetime.fromtimestamp(p.stat().st_mtime, tz=timezone.utc).date()

    wb = load_workbook(p, data_only=True)
    if SHEET not in wb.sheetnames:
        raise BmParseError(f"ABORT: '{SHEET}' not in {p.name}; tabs are {wb.sheetnames}")
    ws = wb[SHEET]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column)]

    h = _find_header(rows)
    months: list[date] = []
    for label in rows[h][3:]:
        m = _parse_month(label)
        if m is None:
            break
        months.append(m)
    if len(months) < MONTH_FLOOR:
        raise BmParseError(
            f"ABORT: '{SHEET}' header row {h + 1} yielded {len(months)} month columns, "
            f"below floor {MONTH_FLOOR}. Labels seen: {rows[h][3:9]!r}"
        )
    if len(set(months)) != len(months):
        dupes = sorted({m.isoformat() for m in months if months.count(m) > 1})
        raise BmParseError(f"ABORT: duplicate month columns in '{SHEET}': {dupes}")

    days = _days_row(rows, h, months)
    blocks, warnings, dupe_keys = _blocks(rows, h, months, days)
    if len(blocks) < BLOCK_FLOOR:
        raise BmParseError(
            f"ABORT: '{SHEET}' yielded {len(blocks)} SKU blocks, below floor {BLOCK_FLOOR}. "
            "A filtered or sorted sheet parses 'successfully' with most rows gone."
        )

    frame, unresolved = _resolve(blocks, months, item_master, aliases)
    placeholders = [b["sku"] for b in blocks if b["placeholder"]]
    if placeholders:
        warnings.append(
            f"BM_PLACEHOLDER_BLOCK: {len(placeholders)} block(s) hold Stores == 1 for every "
            f"month, the sheet's own placeholder shape: {', '.join(sorted(placeholders))}"
        )
    dupe_series = _identical_series(blocks)
    for a, b, metric in dupe_series:
        warnings.append(
            f"BM_IDENTICAL_SERIES: {a} and {b} carry a byte-identical {metric} row across all "
            f"{len(months)} months -- a copy-paste in the sheet, not two independent plans"
        )
    return BmForecast(
        frame=frame,
        months=months,
        snapshot_date=snap,
        source_file=p.name,
        warnings=warnings,
        unresolved=unresolved,
        placeholder_skus=sorted(placeholders),
        duplicate_series=[f"{a}|{b}|{m}" for a, b, m in dupe_series],
    )


def _days_row(rows: list[list[Any]], h: int, months: list[date]) -> list[int]:
    """Locate `Days in Month` by label and verify it against the calendar."""
    for i in range(h + 1, min(h + 8, len(rows))):
        if _norm(rows[i][0]) == DAYS_ROW_LABEL:
            out, bad = [], []
            for j, m in enumerate(months):
                got = _num(rows[i][3 + j], where=f"'{DAYS_ROW_LABEL}' row {i + 1}, {m:%b-%Y}")
                want = calendar.monthrange(m.year, m.month)[1]
                if int(got) != want:
                    bad.append(f"{m:%b-%Y} says {got:g}, calendar says {want}")
                out.append(want)
            if bad:
                raise BmParseError(
                    "ABORT: the sheet's '{}' row disagrees with the calendar on {} month(s): {}. "
                    "It is the denominator of Velocity = Stores x UPSPW x Days/7, so every "
                    "velocity in those columns is wrong.".format(DAYS_ROW_LABEL, len(bad), "; ".join(bad[:5]))
                )
            return out
    raise BmParseError(
        f"ABORT: no '{DAYS_ROW_LABEL}' row within 7 rows of the header in '{SHEET}'. "
        "Velocity = Stores x UPSPW x Days/7 cannot be verified without it."
    )


def _blocks(
    rows: list[list[Any]], h: int, months: list[date], days: list[int]
) -> tuple[list[dict[str, Any]], list[str], dict[str, int]]:
    """One dict per SKU block. `Stores` opens a block: its col B is the SKU and col C the
    unique key; the next row's col B is the description (the sheet's own layout)."""
    blocks: list[dict[str, Any]] = []
    warnings: list[str] = []
    seen_keys: dict[str, int] = {}
    cur: dict[str, Any] | None = None
    for i, r in enumerate(rows):
        label = str(r[0]).strip() if r[0] is not None else None
        if label == "Stores":
            sku = str(r[1]).strip() if r[1] is not None else None
            key = str(r[2]).strip() if r[2] is not None else None
            desc = str(rows[i + 1][1]).strip() if i + 1 < len(rows) and rows[i + 1][1] else None
            if not sku:
                raise BmParseError(f"ABORT: 'Stores' row {i + 1} carries no SKU in column B")
            if not key:
                raise BmParseError(f"ABORT: 'Stores' row {i + 1} ({sku}) carries no Unique Key in column C")
            if key in seen_keys:
                raise BmParseError(
                    f"ABORT: Unique Key {key!r} appears twice, at rows {seen_keys[key]} and "
                    f"{i + 1}. Two blocks for one SKU is a disagreement; it is not resolved here."
                )
            seen_keys[key] = i + 1
            cur = {"row": i + 1, "sku": sku, "key": key, "description": desc, "m": {}}
            blocks.append(cur)
        if cur is not None and label in METRICS:
            cur["m"][label] = [
                _num(r[3 + j], where=f"block {cur['sku']} row {i + 1} ({label}), {m:%b-%Y}")
                for j, m in enumerate(months)
            ]

    for b in blocks:
        missing = [m for m in METRICS if m not in b["m"]]
        if missing:
            raise BmParseError(
                f"ABORT: block {b['sku']!r} (row {b['row']}) is missing metric row(s) {missing}. "
                f"Every block must carry all {len(METRICS)}: {list(METRICS)}"
            )
        # The sheet's own placeholder shape: a single token door wherever the block is
        # live at all. Live 2026-09-09 this catches P-DIS-LGR/BLK/DGR (1 store at UPSPW
        # 12.0 against 7-16 stocked) and K-DIS-2BAB-LGR/PUR (1 against 1,370/1,411).
        nz = [v for v in b["m"]["Stores"] if v]
        b["placeholder"] = bool(nz) and all(v == 1.0 for v in nz)

    warnings += _identities(blocks, months, days)
    return blocks, warnings, seen_keys


def _identities(blocks: list[dict[str, Any]], months: list[date], days: list[int]) -> list[str]:
    """`Total_Demand = Velocity + Load_Orders` aborts (definitional).
    `Velocity = Stores x UPSPW x Days/7` warns (a derivation the author may override)."""
    hard: list[str] = []
    soft: list[str] = []
    for b in blocks:
        for j, m in enumerate(months):
            st, up = b["m"]["Stores"][j], b["m"]["UPSPW"][j]
            ve, lo, td = b["m"]["Velocity"][j], b["m"]["Load_Orders"][j], b["m"]["Total_Demand"][j]
            if abs((ve + lo) - td) > max(0.5, 1e-3 * abs(td)):
                hard.append(f"{b['sku']} {m:%b-%Y}: Velocity {ve:,.1f} + Load_Orders {lo:,.1f} != Total_Demand {td:,.1f}")
            want = st * up * days[j] / 7.0
            if abs(want - ve) > max(0.5, 1e-3 * abs(ve)):
                soft.append(
                    f"{b['sku']} {m:%b-%Y}: Velocity {ve:,.3f} != Stores {st:g} x UPSPW {up:g} "
                    f"x {days[j]}/7 = {want:,.3f} ({100 * (ve - want) / want:+.1f}%)"
                )
    if hard:
        raise BmParseError(
            "ABORT: the sheet's definitional identity Total_Demand = Velocity + Load_Orders "
            f"fails on {len(hard)} cell(s): {'; '.join(hard[:5])}"
        )
    if soft:
        return [
            "BM_VELOCITY_IDENTITY: Velocity != Stores x UPSPW x Days/7 on "
            f"{len(soft)} cell(s) -- a hand override or a copy-paste in the sheet, "
            f"reported not corrected: {'; '.join(soft[:5])}"
        ]
    return []


def _identical_series(blocks: list[dict[str, Any]]) -> list[tuple[str, str, str]]:
    """Two blocks carrying a byte-identical Load_Orders or Velocity row: a copy-paste."""
    out: list[tuple[str, str, str]] = []
    for metric in ("Load_Orders", "Velocity"):
        seen: dict[tuple[float, ...], str] = {}
        for b in blocks:
            vals = tuple(b["m"][metric])
            if not any(vals):
                continue
            prev = seen.get(vals)
            if prev is not None:
                out.append((prev, b["sku"], metric))
            else:
                seen[vals] = b["sku"]
    return out


def _resolve(
    blocks: list[dict[str, Any]],
    months: list[date],
    item_master: ItemMaster | None,
    aliases: AliasMap | None,
) -> tuple[pd.DataFrame, list[dict[str, Any]]]:
    """SKU -> TCIN over `biom_sku`, `rdz_item` and `vendor_style`, aliases applied first.
    Unresolved blocks keep their rows with `tcin` NULL and are listed, never dropped."""
    im = item_master if item_master is not None else ItemMaster.load()
    al = aliases if aliases is not None else AliasMap.load()
    cols = ("biom_sku", "rdz_item", "vendor_style")
    records: list[dict[str, Any]] = []
    unresolved: list[dict[str, Any]] = []
    for b in blocks:
        canon = al.resolve(b["sku"], kind="sku").canonical_sku if al is not None else b["sku"]
        hits = im.tcins_for(canon, columns=cols) or im.tcins_for(b["sku"], columns=cols)
        tcin: int | None = None
        if len(hits) == 1:
            tcin = hits[0]
        elif len(hits) > 1:
            unresolved.append(
                {
                    "bm_sku": b["sku"],
                    "unique_key": b["key"],
                    "description": b["description"],
                    "reason": f"BM_SKU_AMBIGUOUS: matches {len(hits)} TCINs {hits}",
                }
            )
        else:
            unresolved.append(
                {
                    "bm_sku": b["sku"],
                    "unique_key": b["key"],
                    "description": b["description"],
                    "reason": "BM_SKU_NO_TCIN: no item-master row on biom_sku, rdz_item or vendor_style",
                }
            )
        for j, m in enumerate(months):
            rec: dict[str, Any] = {
                "bm_sku": b["sku"],
                "unique_key": b["key"],
                "description": b["description"],
                "tcin": tcin,
                "month_start": pd.Timestamp(m),
                "bm_placeholder": bool(b["placeholder"]),
            }
            for metric in METRICS:
                rec[f"bm_{metric.lower()}"] = b["m"][metric][j]
            records.append(rec)
    frame = pd.DataFrame.from_records(records)
    frame["tcin"] = pd.to_numeric(frame["tcin"], errors="coerce").astype("Int64")
    return frame, unresolved
