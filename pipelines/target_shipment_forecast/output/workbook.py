"""openpyxl workbook writer: README sheet + one sheet per DataFrame.

Numbers are written as numbers (never strings), the header row is frozen and
bold, dates get `yyyy-mm-dd`, integers `#,##0`, floats `#,##0.00`, and column
widths follow the content (capped). Sheet names are sanitised to Excel's rules.
"""

from __future__ import annotations

import math
import re
from collections.abc import Mapping
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
MAX_SHEET_NAME = 31
MAX_COL_WIDTH = 60
MIN_COL_WIDTH = 8


def safe_sheet_name(name: str, taken: set[str]) -> str:
    """Excel-safe, <= 31 chars, unique within the workbook."""
    base = _BAD_SHEET_CHARS.sub("_", str(name)).strip() or "Sheet"
    base = base[:MAX_SHEET_NAME]
    cand, i = base, 2
    while cand in taken:
        suffix = f"_{i}"
        cand = base[: MAX_SHEET_NAME - len(suffix)] + suffix
        i += 1
    taken.add(cand)
    return cand


def _cell_value(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return None if pd.isna(v) else v.to_pydatetime()
    if hasattr(v, "item") and not isinstance(v, str | bytes):  # numpy scalar
        try:
            return _cell_value(v.item())
        except Exception:
            return str(v)
    if isinstance(v, str | int | float | bool | date | datetime):
        return v
    return str(v)


def _number_format(series: pd.Series) -> str | None:
    if pd.api.types.is_bool_dtype(series):
        return None
    if pd.api.types.is_datetime64_any_dtype(series):
        return "yyyy-mm-dd"
    if pd.api.types.is_integer_dtype(series):
        return "#,##0"
    if pd.api.types.is_float_dtype(series):
        nn = series.dropna()
        if len(nn) and (nn == nn.round()).all():
            return "#,##0"
        return "#,##0.00"
    if series.dtype == object:
        nn = series.dropna()
        if len(nn) and all(isinstance(x, date | datetime) for x in nn):
            return "yyyy-mm-dd"
    return None


def _write_frame(ws: Any, df: pd.DataFrame) -> None:
    ws.append([str(c) for c in df.columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    formats = {i + 1: _number_format(df[c]) for i, c in enumerate(df.columns)}
    for row in df.itertuples(index=False, name=None):
        ws.append([_cell_value(v) for v in row])
    for col_idx, fmt in formats.items():
        if fmt:
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=col_idx).number_format = fmt
    ws.freeze_panes = "A2"
    for i, c in enumerate(df.columns, start=1):
        sample = [len(str(v)) for v in df[c].head(200).tolist() if _cell_value(v) is not None]
        width = max([len(str(c)), *sample])
        ws.column_dimensions[get_column_letter(i)].width = max(
            MIN_COL_WIDTH, min(MAX_COL_WIDTH, width + 2)
        )
    ws.auto_filter.ref = ws.dimensions


def _write_readme(ws: Any, readme: Mapping[str, Any]) -> None:
    ws.append(["key", "value"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for k, v in readme.items():
        if isinstance(v, dict | list | tuple):
            v = (
                "; ".join(f"{a}={b}" for a, b in v.items())
                if isinstance(v, dict)
                else ", ".join(map(str, v))
            )
        ws.append([str(k), _cell_value(v)])
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 100
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")


def write_workbook(
    path: str | Path, sheets: Mapping[str, pd.DataFrame], readme: Mapping[str, Any]
) -> Path:
    """Write `README` first, then each sheet in the given order. Returns the path."""
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    taken: set[str] = set()
    ws = wb.active
    ws.title = safe_sheet_name("README", taken)
    _write_readme(ws, readme)
    for name, df in sheets.items():
        w = wb.create_sheet(title=safe_sheet_name(name, taken))
        _write_frame(
            w, df.reset_index(drop=True) if isinstance(df, pd.DataFrame) else pd.DataFrame(df)
        )
    wb.save(out)
    return out
