from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from shipcast.output.workbook import safe_sheet_name, write_workbook


def test_write_workbook(tmp_path: Path) -> None:
    df = pd.DataFrame(
        {
            "tcin": [89854821, 89854823],
            "week": pd.to_datetime(["2026-09-06", "2026-09-06"]),
            "units": [1200.0, 480.0],
            "wape": [0.123, float("nan")],
            "note": ["a", None],
        }
    )
    out = write_workbook(
        tmp_path / "t.xlsx",
        {"Forecast": df, "Bad/Name:Here": df},
        {"as_of": "2026-09-03", "sources": {"plan": "2026-08-29"}},
    )
    wb = load_workbook(out)
    assert wb.sheetnames == ["README", "Forecast", "Bad_Name_Here"]
    ws = wb["Forecast"]
    assert ws.freeze_panes == "A2"
    assert ws["A1"].value == "tcin" and ws["A1"].font.bold
    assert ws["A2"].value == 89854821 and isinstance(ws["A2"].value, int)
    assert ws["C2"].number_format == "#,##0"
    assert ws["D2"].number_format == "#,##0.00"
    assert ws["D3"].value is None
    assert ws["B2"].number_format == "yyyy-mm-dd"
    assert wb["README"]["A2"].value == "as_of" and wb["README"]["B3"].value == "plan=2026-08-29"


def test_safe_sheet_name() -> None:
    taken: set[str] = set()
    assert safe_sheet_name("x" * 40, taken) == "x" * 31
    assert safe_sheet_name("x" * 40, taken) == "x" * 29 + "_2"
